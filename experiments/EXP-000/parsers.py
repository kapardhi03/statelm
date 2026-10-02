"""Readers for the operator's exports: CSV, XLSX and JSON, plus the column mapping file.

Column names vary between exports, so the operator writes a `columns.yaml` naming which source
field holds each canonical field. Nothing here guesses: an unmapped required field is an error,
because silently reading the wrong column would corrupt the data without any visible symptom.
"""

from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from pathlib import Path

REQUIRED_FIELDS = ("conversation_id", "speaker", "timestamp", "text")
OPTIONAL_FIELDS = ("role",)
SUPPORTED_SUFFIXES = {".csv", ".xlsx", ".json"}

#: Sidecars that live beside an export but are not one. extract.py writes
#: `<output>.meta.json` next to its CSV, in the same `data/raw/` tree the scrubber scans, so
#: without this the scrubber would try to read the run's own counts as a conversation export.
IGNORED_NAME_SUFFIXES = (".meta.json",)


class ColumnMapError(Exception):
    pass


@dataclass(frozen=True)
class Message:
    source_conversation_id: str
    speaker: str
    timestamp_raw: object
    text: str
    role: str | None
    source_file: str
    source_row: int


def load_column_map(path: str | Path) -> dict[str, str]:
    """Load columns.yaml. Uses yaml.safe_load, never yaml.load."""
    try:
        import yaml
    except ModuleNotFoundError as exc:  # pragma: no cover - dependency is declared
        raise ColumnMapError(
            "PyYAML is required to read columns.yaml; install this project's dependencies"
        ) from exc
    raw = yaml.safe_load(Path(path).read_text(encoding="utf-8")) or {}
    if not isinstance(raw, dict):
        raise ColumnMapError(f"{path}: expected a mapping of canonical field -> column name")
    mapping = {str(k): str(v) for k, v in raw.items() if v is not None}
    missing = [f for f in REQUIRED_FIELDS if f not in mapping]
    if missing:
        raise ColumnMapError(f"{path}: missing required field(s) {missing}")
    unknown = set(mapping) - set(REQUIRED_FIELDS) - set(OPTIONAL_FIELDS)
    if unknown:
        raise ColumnMapError(f"{path}: unknown field(s) {sorted(unknown)}")
    return mapping


def _row_to_message(row: dict, mapping: dict[str, str], source: str, index: int) -> Message:
    def get(field: str):
        column = mapping.get(field)
        if column is None:
            return None
        if column not in row:
            raise ColumnMapError(
                f"{source}: column {column!r} (mapped to {field!r}) is not present; "
                f"available: {sorted(row)[:12]}"
            )
        return row[column]

    return Message(
        source_conversation_id=str(get("conversation_id") or "").strip(),
        speaker=str(get("speaker") or "").strip(),
        timestamp_raw=get("timestamp"),
        text="" if get("text") is None else str(get("text")),
        role=(str(get("role")).strip() if get("role") not in (None, "") else None),
        source_file=source,
        source_row=index,
    )


def read_csv(path: Path, mapping: dict[str, str]) -> list[Message]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        return [
            _row_to_message(row, mapping, path.name, i)
            for i, row in enumerate(csv.DictReader(fh))
        ]


def read_xlsx(path: Path, mapping: dict[str, str]) -> list[Message]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:  # pragma: no cover - dependency is declared
        raise ColumnMapError(
            "openpyxl is required to read .xlsx; install this project's dependencies"
        ) from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = [str(c).strip() if c is not None else "" for c in next(rows, [])]
        out = []
        for i, values in enumerate(rows):
            if all(v is None for v in values):
                continue
            out.append(_row_to_message(dict(zip(header, values)), mapping, path.name, i))
        return out
    finally:
        workbook.close()


def read_json(path: Path, mapping: dict[str, str]) -> list[Message]:
    """A JSON list of flat objects, or an object wrapping one under a single list key."""
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, dict):
        lists = [v for v in data.values() if isinstance(v, list)]
        if len(lists) != 1:
            raise ColumnMapError(
                f"{path.name}: expected a JSON list of messages, or an object with exactly one "
                f"list value; found {len(lists)} list values"
            )
        data = lists[0]
    if not isinstance(data, list):
        raise ColumnMapError(f"{path.name}: expected a JSON list of messages")
    return [
        _row_to_message(row, mapping, path.name, i)
        for i, row in enumerate(data)
        if isinstance(row, dict)
    ]


READERS = {".csv": read_csv, ".xlsx": read_xlsx, ".json": read_json}


def discover(input_dir: str | Path) -> list[Path]:
    root = Path(input_dir).expanduser().resolve()
    if not root.is_dir():
        raise ColumnMapError(f"{root} is not a directory")
    return sorted(
        p for p in root.rglob("*")
        if p.is_file()
        and p.suffix.lower() in SUPPORTED_SUFFIXES
        and not p.name.startswith("~$")
        and not p.name.lower().endswith(IGNORED_NAME_SUFFIXES)
    )


def read_any(path: Path, mapping: dict[str, str]) -> list[Message]:
    return READERS[path.suffix.lower()](path, mapping)
